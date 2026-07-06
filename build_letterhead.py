#!/usr/bin/env python3
"""Generate a modern legal letterhead as a Word (.docx) file.

The masthead lives in the document header and the contact details in the
footer, so both repeat on every page while the body stays free for typing.
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---- palette (matches the HTML letterhead) ----------------------------------
INK   = RGBColor(0x1B, 0x2A, 0x41)   # deep navy ink
SLATE = RGBColor(0x5A, 0x64, 0x78)   # secondary text
BRASS = RGBColor(0x9C, 0x7C, 0x46)   # restrained accent
BODY  = RGBColor(0x26, 0x31, 0x3F)

SERIF = "Palatino Linotype"
SANS  = "Calibri"


def set_font(run, name):
    """Apply a font family for both Latin and complex-script ranges."""
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), name)


# CT_RPr child order (subset) — used to insert elements at schema-correct spots.
_RPR_AFTER_CAPS = (
    "w:smallCaps", "w:strike", "w:dstrike", "w:outline", "w:shadow", "w:emboss",
    "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden",
    "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz", "w:szCs",
    "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText",
    "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
)
_RPR_AFTER_SPACING = _RPR_AFTER_CAPS[_RPR_AFTER_CAPS.index("w:w"):]
# CT_PPr children that follow pBdr.
_PPR_AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def set_tracking(run, twentieths):
    """Letter-spacing in twentieths of a point (w:spacing)."""
    rpr = run._element.get_or_add_rPr()
    spc = OxmlElement("w:spacing")
    spc.set(qn("w:val"), str(twentieths))
    rpr.insert_element_before(spc, *_RPR_AFTER_SPACING)


def set_caps(run):
    rpr = run._element.get_or_add_rPr()
    caps = OxmlElement("w:caps")
    caps.set(qn("w:val"), "true")
    rpr.insert_element_before(caps, *_RPR_AFTER_CAPS)


def _pbdr(paragraph):
    """Return this paragraph's pBdr, creating it at the right position."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.insert_element_before(borders, *_PPR_AFTER_PBDR)
    return borders


def bottom_border(paragraph, color, size=12, space=6):
    """Draw a rule beneath a paragraph. size = eighths of a point."""
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    _pbdr(paragraph).append(bottom)


def top_border(paragraph, color, size=4, space=8):
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(size))
    top.set(qn("w:space"), str(space))
    top.set(qn("w:color"), color)
    _pbdr(paragraph).append(top)


def cell_borders(cell, color, size=6):
    """Thin border on all sides of a table cell (for the monogram box)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def no_space(paragraph, before=0, after=0, line=None):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line


# ---- document ---------------------------------------------------------------
doc = Document()

section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1.9)      # room for the header masthead
section.bottom_margin = Inches(1.1)   # room for the footer band
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.6)
section.footer_distance = Inches(0.5)

# Base "Normal" style
normal = doc.styles["Normal"]
normal.font.name = SERIF
normal.font.size = Pt(11)
normal.font.color.rgb = BODY

# ============================ HEADER (masthead) ==============================
header = section.header
header.is_linked_to_previous = False
# clear default empty paragraph
header_para0 = header.paragraphs[0]
header_para0.text = ""

# A 2-column table: [ monogram box ] [ name + title ]
htable = header.add_table(rows=1, cols=2, width=Inches(6.5))
htable.autofit = False
htable.allow_autofit = False
mono_cell, id_cell = htable.rows[0].cells
mono_cell.width = Inches(0.85)
id_cell.width = Inches(5.65)
mono_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
id_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# monogram
cell_borders(mono_cell, "9C7C46", size=6)
mp = mono_cell.paragraphs[0]
mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
no_space(mp, before=6, after=6)
mr = mp.add_run("CSR")
set_font(mr, SERIF)
mr.font.size = Pt(16)
mr.font.color.rgb = INK
set_tracking(mr, 20)

# name
np = id_cell.paragraphs[0]
np.alignment = WD_ALIGN_PARAGRAPH.LEFT
no_space(np, before=0, after=4, line=1.0)
nr = np.add_run("Carlos M. Santini Rodríguez")
set_font(nr, SERIF)
nr.font.size = Pt(19)
nr.font.color.rgb = INK
set_caps(nr)
set_tracking(nr, 28)
esq = np.add_run(", Esq.")
set_font(esq, SERIF)
esq.font.size = Pt(19)
esq.font.italic = True
esq.font.color.rgb = BRASS
set_tracking(esq, 4)

# title
tp = id_cell.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.LEFT
no_space(tp, before=0, after=0)
tr = tp.add_run("Attorney & Counselor at Law")
set_font(tr, SANS)
tr.font.size = Pt(9)
tr.font.color.rgb = SLATE
tr.font.bold = True
set_caps(tr)
set_tracking(tr, 68)

# brass rule beneath the masthead
rule_p = header.add_paragraph()
no_space(rule_p, before=8, after=0)
bottom_border(rule_p, "9C7C46", size=12, space=4)

# ============================ FOOTER (contact band) ==========================
footer = section.footer
footer.is_linked_to_previous = False
fp = footer.paragraphs[0]
fp.text = ""
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
no_space(fp, before=6, after=0)
top_border(fp, "E3DED2", size=4, space=8)

segments = [
    ("PO Box 30309", SLATE, False),
    ("San Juan, PR 00929-0544", SLATE, False),
    ("Tel. 787.636.0770", RGBColor(0x3A, 0x46, 0x58), True),
    ("carsantini@gmail.com", SLATE, False),
]
for i, (text, color, strong) in enumerate(segments):
    if i:
        dot = fp.add_run("   ·   ")
        set_font(dot, SANS)
        dot.font.size = Pt(8)
        dot.font.color.rgb = BRASS
        dot.font.bold = True
    run = fp.add_run(text)
    set_font(run, SANS)
    run.font.size = Pt(8)
    run.font.color.rgb = color
    run.font.bold = strong
    if not text[0].islower():   # letter-space the uppercase-style items
        set_caps(run)
        set_tracking(run, 30)

# ============================ BODY (editable sample) =========================
date_p = doc.add_paragraph()
no_space(date_p, before=6, after=18)
dr = date_p.add_run("July 6, 2026")
set_font(dr, SERIF); dr.font.size = Pt(11); dr.font.color.rgb = BODY

recipient = ("[Recipient Name]\n[Title / Organization]\n"
             "[Street Address]\n[City, State ZIP]")
rp = doc.add_paragraph()
no_space(rp, before=0, after=18, line=1.15)
rr = rp.add_run(recipient)
set_font(rr, SERIF); rr.font.size = Pt(11); rr.font.color.rgb = BODY

sal = doc.add_paragraph()
no_space(sal, before=0, after=12)
sr = sal.add_run("Dear [Recipient]:")
set_font(sr, SERIF); sr.font.size = Pt(11); sr.font.color.rgb = BODY

body_paras = [
    "This letterhead is ready for use. Replace this text with the body of your "
    "correspondence—the masthead and footer are set in the page header and "
    "footer, so they repeat automatically on every page.",
    "The document uses standard U.S. Letter dimensions with legal-standard "
    "margins, so it may be used directly for formal correspondence, engagement "
    "letters, and filings.",
]
for text in body_paras:
    p = doc.add_paragraph()
    no_space(p, before=0, after=12, line=1.4)
    r = p.add_run(text)
    set_font(r, SERIF); r.font.size = Pt(11); r.font.color.rgb = BODY

# signature block
close_p = doc.add_paragraph()
no_space(close_p, before=18, after=36)
cr = close_p.add_run("Respectfully,")
set_font(cr, SERIF); cr.font.size = Pt(11); cr.font.color.rgb = BODY

who_p = doc.add_paragraph()
no_space(who_p, before=0, after=2)
wr = who_p.add_run("Carlos M. Santini Rodríguez, Esq.")
set_font(wr, SERIF); wr.font.size = Pt(11); wr.font.bold = True; wr.font.color.rgb = INK

whot_p = doc.add_paragraph()
no_space(whot_p, before=0, after=0)
wtr = whot_p.add_run("Attorney & Counselor at Law")
set_font(wtr, SANS); wtr.font.size = Pt(8.5); wtr.font.color.rgb = SLATE
set_caps(wtr); set_tracking(wtr, 40)

doc.save("/home/user/great_files/letterhead.docx")
print("wrote letterhead.docx")
